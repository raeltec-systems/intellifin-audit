# Synthetic Story 22.5 document fixtures. Every value is invented test data.
# Regenerate inside the pinned general analysis image (it has the libraries):
#   docker run --rm --network none -v "$PWD/fixtures/documents:/out" \
#     127.0.0.1:5000/zobba/analysis-general@<GENERAL_IMAGE_DIGEST> \
#     python3 -I -B /out/generate.py /out
# The generated files are committed; tests rely on their exact bytes.
import datetime
import io
import json
import os
import re
import sys
import zipfile
import zlib

OUT = sys.argv[1]
FIXED = datetime.datetime(2026, 1, 1, 0, 0, 0)


def save(name, data):
    with open(os.path.join(OUT, name), "wb") as handle:
        handle.write(data)


def normalized_zip(data):
    """Rewrite a package with fixed member timestamps so bytes are stable."""
    source = zipfile.ZipFile(io.BytesIO(data))
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            body = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                # Libraries stamp the save time into the package properties.
                body = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
                    rb"\g<1>2026-01-01T00:00:00Z\g<2>",
                    body,
                )
            out.writestr(fixed, body)
    return target.getvalue()


def workbook():
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    rows = [
        ("Account", "Balance"),
        ("Cash", 830.5),
        ("Receivables", 419.5),
        ("Total", 1250),
    ]
    for row in rows:
        sheet.append(row)
    for row in range(2, 5):
        sheet.cell(row=row, column=2).number_format = "#,##0.00"
    book.properties.created = FIXED
    book.properties.modified = FIXED
    buffer = io.BytesIO()
    book.save(buffer)
    return normalized_zip(buffer.getvalue())


def ap_invoices():
    """Story 24.1 AP invoices: amounts with a display format, a subject with
    leading zeros, a day-first date and an instant with its source offset."""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    rows = [
        ("Invoice", "Amount", "Currency", "Approver", "Invoice date", "Approved at", "In approval history"),
        ("INV-0042", 1250, "USD", "J. Smith", "05/03/2026", "2026-03-05T23:30:00-05:00", "Yes"),
        ("000123", 980, "USD", "J. Smith", "06/03/2026", "2026-03-06T09:00:00+00:00", "Yes"),
        ("INV-0054", 1250, "USD", "", "07/03/2026", "", "No"),
        ("INV-0055", 4200, "EUR", "A. Jones", "08/03/2026", "2026-03-08T10:00:00+01:00", "Yes"),
        # Story 24.1 review: one row per matrix invoice, so each fact's
        # subject and currency stand in the same row as its value.
        ("INV-0050", 1250, "USD", "J. Smith", "09/03/2026", "", "Yes"),
        ("INV-0051", 1250, "USD", "", "10/03/2026", "", ""),
        ("INV-0052", 1250, "USD", "", "11/03/2026", "", ""),
        ("INV-0053", 1250, "USD", "J. Smith", "12/03/2026", "", ""),
        ("INV-0057", 3000, "EUR", "A. Jones", "13/03/2026", "", "Yes"),
    ]
    for row in rows:
        sheet.append(row)
    for row in range(2, 11):
        sheet.cell(row=row, column=2).number_format = "#,##0.00"
    book.properties.created = FIXED
    book.properties.modified = FIXED
    buffer = io.BytesIO()
    book.save(buffer)
    return normalized_zip(buffer.getvalue())


def contract_docx():
    import docx

    document = docx.Document()
    document.add_paragraph("Services agreement between Example Supplier and Example Client.")
    document.add_paragraph("Payment is due within 30 days of the invoice date.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Fee"
    table.cell(0, 1).text = "1,250.00"
    table.cell(1, 0).text = "Term"
    table.cell(1, 1).text = "12 months"
    document.core_properties.created = FIXED
    document.core_properties.modified = FIXED
    document.core_properties.last_printed = FIXED
    buffer = io.BytesIO()
    document.save(buffer)
    return normalized_zip(buffer.getvalue())


def pdf(pages):
    """Minimal hand-written PDF. Each page is (text_lines, with_image)."""
    objects = []

    def add(body):
        objects.append(body)
        return len(objects)

    catalog = add(None)
    pages_id = add(None)
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for lines, with_image in pages:
        stream = b""
        resources = b"<< /Font << /F1 %d 0 R >>" % font
        if with_image:
            # 8x8 grey image: a "scan" with no text layer at all.
            pixels = zlib.compress(bytes([128] * 8 * 8 * 3))
            image = add(
                b"<< /Type /XObject /Subtype /Image /Width 8 /Height 8 /ColorSpace /DeviceRGB "
                b"/BitsPerComponent 8 /Filter /FlateDecode /Length %d >>\nstream\n%s\nendstream"
                % (len(pixels), pixels)
            )
            resources += b" /XObject << /Im1 %d 0 R >>" % image
            stream += b"q 400 0 0 500 100 150 cm /Im1 Do Q\n"
        resources += b" >>"
        for index, line in enumerate(lines):
            stream += b"BT /F1 14 Tf 72 %d Td (%s) Tj ET\n" % (720 - index * 24, line.encode("latin-1"))
        content = add(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
        kids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Resources %s /Contents %d 0 R >>"
                % (pages_id, resources, content)
            )
        )
    objects[catalog - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id
    objects[pages_id - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids),
        len(kids),
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref))
    return out.getvalue()


EMAIL = (
    "From: Example Supplier <supplier@example.test>\r\n"
    "To: Example Client <client@example.test>\r\n"
    "Subject: Payment terms\r\n"
    "Date: Thu, 1 Jan 2026 00:00:00 +0000\r\n"
    "MIME-Version: 1.0\r\n"
    'Content-Type: multipart/alternative; boundary="b1"\r\n'
    "\r\n"
    "--b1\r\n"
    "Content-Type: text/plain; charset=utf-8\r\n"
    "\r\n"
    "We confirm payment within 30 days.\r\n"
    "--b1\r\n"
    "Content-Type: text/html; charset=utf-8\r\n"
    "\r\n"
    "<p>We confirm payment within <b>30 days</b>.</p>\r\n"
    "--b1--\r\n"
)


def zip_bomb():
    """A package declaring far more expansion than any bound allows."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for name in ("[Content_Types].xml", "xl/workbook.xml"):
            out.writestr(zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0)), b"<x/>")
        out.writestr(
            zipfile.ZipInfo("xl/worksheets/sheet1.xml", date_time=(2026, 1, 1, 0, 0, 0)),
            b"\0" * (64 * 1024 * 1024),
            compress_type=zipfile.ZIP_DEFLATED,
        )
    return buffer.getvalue()


def macro_workbook():
    """An .xlsm-shaped package carrying a VBA project; it must never run."""
    data = workbook()
    source = zipfile.ZipFile(io.BytesIO(data))
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            out.writestr(zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0)), source.read(info.filename))
        out.writestr(
            zipfile.ZipInfo("xl/vbaProject.bin", date_time=(2026, 1, 1, 0, 0, 0)),
            b"Attribute VB_Name = \"Auto_Open\"\nShell \"curl http://169.254.169.254\"\n",
        )
    return target.getvalue()


# A 64x32 baseline JPEG "photo" (a drawn frame, no text), encoded once.
PHOTO = bytes.fromhex(
    "ffd8ffe000104a46494600010100000100010000ffdb0043000d090a0b0a080d0b0a0b0e0e0d"
    "0f13201513121213271c1e17202e2931302e292d2c333a4a3e333646372c2d405741464c4e52"
    "5352323e5a615a50604a51524fffdb0043010e0e0e131113261515264f352d354f4f4f4f4f4f"
    "4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f4f"
    "4f4f4f4f4f4fffc00011080020004003012200021101031101ffc4001f000001050101010101"
    "0100000000000000000102030405060708090a0bffc400b51000020103030204030505040400"
    "00017d01020300041105122131410613516107227114328191a1082342b1c11552d1f0243362"
    "7282090a161718191a25262728292a3435363738393a434445464748494a535455565758595a"
    "636465666768696a737475767778797a838485868788898a92939495969798999aa2a3a4a5a6"
    "a7a8a9aab2b3b4b5b6b7b8b9bac2c3c4c5c6c7c8c9cad2d3d4d5d6d7d8d9dae1e2e3e4e5e6e7"
    "e8e9eaf1f2f3f4f5f6f7f8f9faffc4001f010003010101010101010101000000000000010203"
    "0405060708090a0bffc400b51100020102040403040705040400010277000102031104052131"
    "061241510761711322328108144291a1b1c109233352f0156272d10a162434e125f11718191a"
    "262728292a35363738393a434445464748494a535455565758595a636465666768696a737475"
    "767778797a82838485868788898a92939495969798999aa2a3a4a5a6a7a8a9aab2b3b4b5b6b7"
    "b8b9bac2c3c4c5c6c7c8c9cad2d3d4d5d6d7d8d9dae2e3e4e5e6e7e8e9eaf2f3f4f5f6f7f8f9"
    "faffda000c03010002110311003f00edafaf2ea2bdb7b4b3b68669258e490996631850a50765"
    "6c9f9c7e54df3b5aff00a07e9fff0081cfff00c6a89ffe464b2ffaf3b8ff00d0e1ad1a00cef3"
    "b5affa07e9ff00f81cff00fc6a8f3b5aff00a07e9fff0081cfff00c6ab468a00cef3b5affa07"
    "e9ff00f81cff00fc6a8f3b5aff00a07e9fff0081cfff00c6ab468a00cef3b5affa07e9ff00f8"
    "1cff00fc6a9d6379752dedc5a5e5b430c91471c80c531903062e3baae0fc87f3abf59d07fc8c"
    "97bff5e76fff00a1cd4004ff00f232597fd79dc7fe870d68d50beb3ba96f6deeecee618648a3"
    "92322584c8183143d99707e41f9d37c9d6bfe821a7ff00e00bff00f1da00d1a2b3bc9d6bfe82"
    "1a7ffe00bfff001da3c9d6bfe821a7ff00e00bff00f1da00d1a2b3bc9d6bfe821a7ffe00bfff"
    "001da3c9d6bfe821a7ff00e00bff00f1da00d1ace83fe464bdff00af3b7ffd0e6a3c9d6bfe82"
    "1a7ffe00bfff001da758d9dd457b717779730cd24b1c7181142630a14b9eecd93f39fca803ff"
    "d9"
)


POPULATION_HEADERS = ("Invoice", "Invoice date", "Amount", "Approved", "Approved before payment", "Currency")


def ap_population_rows():
    """Story 24.2 AP population: 1,250 invoices and 6 exact duplicate rows.
    1-50 are dated in 2025 (outside the 2026 basis period); 51-250 are
    500.00 USD; 251-1250 are 1,500.00 USD. Approved is No for 251-265 and
    Pending for 276-285; Approved before payment is No for 259-270 and
    Unknown for 286-295. Rows of 1000-1005 are repeated exactly at the end."""
    rows = []
    for i in range(1, 1251):
        approved = "No" if 251 <= i <= 265 else "Pending" if 276 <= i <= 285 else "Yes"
        before = "No" if 259 <= i <= 270 else "Unknown" if 286 <= i <= 295 else "Yes"
        rows.append((
            "INV-%05d" % i,
            "2025-12-15" if i <= 50 else "2026-03-05",
            500 if i <= 250 else 1500,
            approved,
            before,
            "USD",
            "note %d" % i,
        ))
    rows += rows[999:1005]
    return rows


def ap_population():
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Invoices"
    sheet.append(("AP invoice export (synthetic)",))
    sheet.append(())
    sheet.append(POPULATION_HEADERS + ("Note",))
    for row in ap_population_rows():
        sheet.append(row)
    for row in range(4, sheet.max_row + 1):
        sheet.cell(row=row, column=3).number_format = "#,##0.00"
    summary = book.create_sheet("Summary")
    summary.append(("Invoice count", 1250))
    summary.append(("Invoice count (cover sheet)", 1251))
    for row in (1, 2):
        summary.cell(row=row, column=2).number_format = "#,##0"
    low = book.create_sheet("Low")
    low.append(POPULATION_HEADERS)
    for i in range(1, 21):
        low.append(("LOW-%03d" % i, "2026-02-01", 1000, "Yes", "Yes", "USD"))
    filtered = book.create_sheet("Filtered")
    filtered.append(POPULATION_HEADERS)
    mismatch = book.create_sheet("Mismatch")
    mismatch.append(("Invoice", "Invoice date", "Amount ", "Approved", "Approved before payment", "Currency"))
    mismatch.append(("INV-1", "2026-02-01", 1500, "Yes", "Yes", "USD"))
    edge = book.create_sheet("Edge")
    edge.append(POPULATION_HEADERS)
    edge.append(("000123", "05/03/2026", 1500, "Yes", "Yes", "USD"))
    edge.append(("000123", "05/03/2026", 1600, "Yes", "Yes", "USD"))
    edge.append(("INV-EDGE-1", "31/02/2026", 1500, "Yes", "Yes", "USD"))
    edge.append(("INV-EDGE-2", None, 1500, "Yes", "Yes", "USD"))
    edge.append((None, "05/03/2026", 1500, "Yes", "Yes", "USD"))
    edge.append(("INV-EDGE-3", "05/03/2026", 1500, "Yes", "Yes", "USD"))
    lines = book.create_sheet("Lines")
    lines.append(("Invoice", "Line", "Invoice date", "Amount", "Approved", "Approved before payment", "Currency"))
    for invoice in ("INV-00251", "INV-00300", "INV-00301"):
        for line in ("1", "2"):
            lines.append((invoice, line, "2026-03-05", 1500, "Yes", "Yes", "USD"))
    book.properties.created = FIXED
    book.properties.modified = FIXED
    buffer = io.BytesIO()
    book.save(buffer)
    return normalized_zip(buffer.getvalue())


def ap_population_csv():
    rows = ['Invoice,Invoice date,Amount,Approved,Approved before payment,Currency,Note']
    rows.append('INV-C1,2026-02-01,"1,500.00",Yes,Yes,USD,"a note, with a comma"')
    rows.append('INV-C2,2026-02-02,"1,500.00",No,Yes,USD,"say ""hi"""')
    rows.append('INV-C3,2025-12-31,"1,500.00",Yes,Yes,USD,')
    rows.append('INV-C4,2026-02-03,980.00,Yes,Yes,USD,')
    return ("\r\n".join(rows) + "\r\n").encode("utf-8")


def ap_population_expected():
    """Every count the AC1/AC2 scenarios must reproduce, computed here from
    the rows above, independently of the Rust code under test."""
    rows = ap_population_rows()
    keys = [r[0] for r in rows]
    subjects = sorted(set(keys))
    first = {}
    for r in rows:
        first.setdefault(r[0], r)
    in_scope = [k for k in subjects if first[k][1].startswith("2026")]
    applicable = [k for k in in_scope if first[k][2] > 1000]
    a_no = [k for k in applicable if first[k][3] == "No"]
    a_und = [k for k in applicable if first[k][3] == "Pending"]
    b_no = [k for k in applicable if first[k][4] == "No"]
    b_und = [k for k in applicable if first[k][4] == "Unknown"]
    failing = sorted(set(a_no) | set(b_no))
    unresolved = sorted((set(a_und) | set(b_und)) - set(failing))
    return {
        "received_rows": len(rows),
        "duplicate_rows": len(rows) - len(subjects),
        "distinct_subjects": len(subjects),
        "excluded_out_of_period": len(subjects) - len(in_scope),
        "in_scope": len(in_scope),
        "applicable": len(applicable),
        "not_applicable": len(in_scope) - len(applicable),
        "a": {"met": len(applicable) - len(a_no) - len(a_und), "not_met": len(a_no), "undetermined": len(a_und)},
        "b": {"met": len(applicable) - len(b_no) - len(b_und), "not_met": len(b_no), "undetermined": len(b_und)},
        "failing": len(failing),
        "failing_opportunities": len(a_no) + len(b_no),
        "unresolved": len(unresolved),
        "meets_every_applicable": len(applicable) - len(failing) - len(unresolved),
        "not_applicable_to_any": len(in_scope) - len(applicable),
        "conclusion": "exceptions_found_inconclusive",
    }


save("Receipt photo.jpg", PHOTO)
save("Trial balance.xlsx", workbook())
save("Contract.docx", contract_docx())
save(
    "Contract.pdf",
    pdf(
        [
            (["Services agreement", "Example Supplier and Example Client"], False),
            (["Payment terms", "Payment is due within 30 days."], False),
            (["Signatures", "Signed by A. Auditor"], False),
        ]
    ),
)
save("Scanned letter.pdf", pdf([([], True), ([], True), ([], True)]))
save("Payment terms.eml", EMAIL.encode("utf-8"))
save("Zip bomb.xlsx", zip_bomb())
save("Encrypted.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504)
save("Macro workbook.xlsm", macro_workbook())
save("Notes.txt", "Opening balance 1,250.00 agreed to the ledger.\n".encode("utf-8"))
save("AP invoices.xlsx", ap_invoices())
save(
    "AP policy.txt",
    "Invoices above 1,000.00 USD require an approver. No invoice may exceed 5,000.00 USD.\n".encode("utf-8"),
)
save("AP population.xlsx", ap_population())
save("AP population.csv", ap_population_csv())
save("AP population expected.json", (json.dumps(ap_population_expected(), indent=2, sort_keys=True) + "\n").encode("utf-8"))
